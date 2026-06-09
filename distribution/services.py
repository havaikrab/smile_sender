import logging
import os
import time
from io import BytesIO

import openpyxl
from django.contrib.auth.models import AbstractBaseUser, AnonymousUser
from django.core.exceptions import ObjectDoesNotExist, ValidationError
from django.core.files.base import File
from django.core.mail import send_mail
from django.db.models import QuerySet
from django.forms import BaseForm
from django.http import FileResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from config.settings import EMAIL_HOST_USER
from distribution.models import Attempt, Mailing, Recipient
from users.models import CustomUser

distribution_logger = logging.getLogger("distribution_logger")
os.makedirs("logs/", exist_ok=True)
file_handler = logging.FileHandler("logs/distribution_logger.log", mode="a", encoding="utf-8")
file_formatter = logging.Formatter("%(asctime)s %(levelname)s: %(message)s")
file_handler.setFormatter(file_formatter)
distribution_logger.addHandler(file_handler)
distribution_logger.setLevel(logging.INFO)


class ExcelManager:
    """Класс работы с Excel-файлами"""

    __columns = {
        "A": "email",
        "B": "first_name",
        "C": "middle_name",
        "D": "last_name",
        "E": "comment",
    }

    def __init__(self, form: BaseForm, excel_file: File):
        """Создание менеджера для загружаемого excel-файла"""

        self.__file = excel_file
        self.__form = form
        self.__sheet = Workbook().active
        self.report: dict = {"success_operations": list(), "existing_objects": list(), "invalid_rows": list()}

    def check_file(self) -> None:
        """Проверка корректности структуры загружаемого файла"""

        data = openpyxl.load_workbook(self.__file)
        if "recipients" not in data.sheetnames:
            self.__form.add_error("excel_file", 'Файл должен содержать лист "recipients"')
            raise KeyError
        sheet = data["recipients"]
        for k, v in ExcelManager.__columns.items():
            if sheet[f"{k}1"].value != v:
                self.__form.add_error("excel_file", 'Лист "recipients" содержит некорректные названия столбцов')
                raise KeyError
        self.__sheet = sheet

    def create_recipients(self) -> None:
        """Запись информации о получателях рассылок в БД"""

        if isinstance(self.__sheet, Worksheet):
            existing_emails = list(Recipient.objects.all().values_list("email", flat=True))
            uploaded_emails = set([self.__sheet[f"A{row}"].value for row in range(2, self.__sheet.max_row + 1)])
            new_emails = uploaded_emails.difference(existing_emails)
            new_emails.discard(None)
            new_emails.discard("")
            valid_recipients = list()
            for row in range(2, self.__sheet.max_row + 1):
                recipient_dict = {v: self.__sheet[f"{k}{row}"].value for k, v in ExcelManager.__columns.items()}
                row_email = recipient_dict["email"]
                if row_email in new_emails:
                    try:
                        recipient = Recipient(**recipient_dict)
                        recipient.full_clean()
                        valid_recipients.append(recipient)
                        self.report["success_operations"].append(
                            f"Получатель №{row} с email {row_email} добавлен в базу данных"
                        )
                    except ValidationError:
                        self.report["invalid_rows"].append(str(row))
                    new_emails.discard(row_email)
                else:
                    if row_email is not None and row_email != "":
                        self.report["existing_objects"].append(
                            f"Получатель №{row} с email {row_email} уже существует в базе данных"
                        )
            Recipient.objects.bulk_create(valid_recipients, ignore_conflicts=True)

    @staticmethod
    def send_excel_form() -> FileResponse:
        """Подготавливает к отправке стандартную excel-форму списка получателей"""

        new_excel = Workbook()
        current_sheet = new_excel.active
        if isinstance(current_sheet, Worksheet):
            current_sheet.title = "recipients"
            current_sheet.append(("email", "first_name", "middle_name", "last_name", "comment"))
        buffer = BytesIO()
        new_excel.save(buffer)
        buffer.seek(0)

        response = FileResponse(
            buffer,
            as_attachment=True,
            filename="recipients_form.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        return response


def group_context(queryset: QuerySet[Mailing]) -> dict:
    """Сортирует список объектов рассылок в соответствии со статусом"""

    time_now = timezone.now()
    sorted_context: dict[str, list] = {"created": list(), "ready": list(), "started": list(), "completed": list()}
    for mailing in queryset:
        if mailing.end_time < time_now or mailing.status == "completed":
            mailing.status = "completed"
            sorted_context["completed"].append(mailing)
        elif mailing.status == "started":
            sorted_context["started"].append(mailing)
        elif mailing.start_time <= time_now <= mailing.end_time:
            sorted_context["ready"].append(mailing)
        else:
            sorted_context["created"].append(mailing)
        Mailing.objects.bulk_update(sorted_context["completed"], ["status"])
    return sorted_context


def check_readiness(mailing_id: int) -> Mailing:
    """Определяет готовность рассылки к запуску"""

    mailing = Mailing.objects.get(pk=mailing_id)
    if mailing.status != "created":
        raise ValueError(
            f'Запускаемая рассылка должна иметь статус "created". Текущий статус рассылки {mailing.status}.'
        )
    time_now = timezone.now()
    if mailing.start_time > time_now or mailing.end_time < time_now:
        raise ValueError("Рассылка не может быть запущена в настоящее время.")
    return mailing


def send_mails(mailing: Mailing) -> None:
    """Отправка писем с установленным временным интервалом"""

    distribution_logger.info(f"Запущена рассылка id{mailing.pk}.")
    mailing.status = "started"
    mailing.save()
    recipients = mailing.recipients.all()
    success_count = 0
    fail_count = 0
    attempts_list = list()
    time_to_sleep = int(os.getenv("SENDING_INTERVAL", 300))
    for recipient in recipients:
        try:
            send_mail(
                mailing.message.title,
                mailing.message.content,
                from_email=EMAIL_HOST_USER,
                recipient_list=[recipient.email],
                fail_silently=False,
            )
            smtp_response = "250 OK: message accepted for delivery"
            status = "success"
            success_count += 1
        except Exception as exc:
            smtp_response = str(exc)
            status = "fail"
            distribution_logger.warning(f"Ошибка при обращении к SMTP-серверу: {exc}.")
            fail_count += 1
        send_time = timezone.localtime()
        print(send_time)
        attempts_list.append(
            Attempt(
                mailing=mailing, recipient=recipient, status=status, server_response=smtp_response, send_time=send_time
            )
        )
        time.sleep(time_to_sleep)
    Attempt.objects.bulk_create(attempts_list)
    mailing.status = "completed"
    mailing.save()
    distribution_logger.info(
        f"Рассылка id{mailing.pk} завершена. Отправлено - {success_count}, не отправлено - {fail_count} писем."
    )


def execute_mailing(mailing_id: int) -> None:
    """Запускает рассылку"""

    try:
        mailing = check_readiness(mailing_id)
    except ObjectDoesNotExist:
        distribution_logger.error("Попытка запуска несуществующей рассылки.")
    except ValueError as exc:
        distribution_logger.error(f"{exc}")
    else:
        send_mails(mailing)


def get_statistics(user: AbstractBaseUser | AnonymousUser) -> dict:
    """Возвращает словарь статистики для главной страницы приложения"""

    grouped_context = group_context(Mailing.objects.all())
    time_now = timezone.now()
    started_list = grouped_context["started"]
    active_mailing = [mailing for mailing in started_list if mailing.start_time <= time_now <= mailing.end_time]
    grouped_context["active_count"] = len(active_mailing)
    try:
        max_id = Mailing.objects.latest("id").pk
    except ObjectDoesNotExist:
        max_id = 0
    grouped_context["total_mailing"] = max_id
    grouped_context["recipients_count"] = Recipient.objects.count()
    if isinstance(user, CustomUser):
        users_attempts = Attempt.objects.filter(mailing__message__author=user)
        grouped_context["total_attempts"] = len(users_attempts)
        success_attempts = users_attempts.filter(status="success")
        grouped_context["success_attempts"] = len(success_attempts)
        failed_attempts = users_attempts.filter(status="fail")
        grouped_context["failed_attempts"] = len(failed_attempts)
    return grouped_context
